"""
report.py
---------
Convierte la lista de anunciantes en entregables "bonitos":
  - informe HTML autonomo (un solo archivo, abrible en el navegador)
  - Excel (.xlsx) con formato, filtros e hipervinculos
  - CSV (compatibilidad / importacion)
  - JSON (para integraciones)
"""

from __future__ import annotations

import os
import csv
import json
import html
import datetime
from collections import Counter

from core.sectors import normalize_sector

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# --- Paleta Nevo para el Excel (sin '#', formato ARGB de openpyxl) ---
XL_INK = "FF0A0A0A"
XL_ORANGE = "FFFF4B00"
XL_GRAY_50 = "FFF4F4F2"
XL_GRAY_100 = "FFECECEA"
XL_WHITE = "FFFFFFFF"
XL_GREEN = "FF1F7A4D"
XL_AMBER = "FF9A6A10"
XL_RED = "FFB8430F"


def export_csv(advertisers, path: str):
    cols = ["brand", "sector", "pages", "website", "email", "phone",
            "ad_size", "confidence", "method", "notes"]
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Marca", "Sector", "Paginas", "Web", "Email",
                    "Telefono", "Tamano anuncio", "Confianza",
                    "Metodo", "Notas"])
        for a in advertisers:
            d = a if isinstance(a, dict) else a.__dict__
            w.writerow([
                d["brand"], d["sector"],
                ", ".join(str(p) for p in d["pages"]),
                d["website"], d["email"], d["phone"], d["ad_size"],
                f'{d["confidence"]:.0%}', d["method"], d["notes"],
            ])


def export_xlsx(advertisers, path: str, meta: dict):
    """Genera un Excel profesional con dos hojas: Resumen y Anunciantes.

    Hoja Anunciantes: cabecera negra estilo Nevo, autofiltro, fila fija,
    hipervinculos en web/email, confianza coloreada. Pensado para que el
    cliente filtre y trabaje directamente.
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError("openpyxl no esta instalado (pip install openpyxl)")

    wb = openpyxl.Workbook()

    thin = Side(style="thin", color=XL_GRAY_100)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name="Arial", bold=True, color=XL_WHITE, size=10)
    head_fill = PatternFill("solid", fgColor=XL_INK)
    cell_font = Font(name="Arial", size=10, color=XL_INK)
    link_font = Font(name="Arial", size=10, color=XL_INK, underline="single")

    # ---------------- Hoja Resumen ----------------
    ws0 = wb.active
    ws0.title = "Resumen"
    ws0.sheet_view.showGridLines = False

    ws0["B2"] = "NEVO"
    ws0["B2"].font = Font(name="Arial Black", bold=True, size=22,
                          color=XL_INK)
    ws0["B3"] = "INFORME DE ANUNCIANTES"
    ws0["B3"].font = Font(name="Arial", size=10, color=XL_ORANGE, bold=True)

    n = len(advertisers)
    con_web = sum(1 for a in advertisers
                  if (a if isinstance(a, dict) else a.__dict__).get("website"))
    con_contacto = sum(
        1 for a in advertisers
        if any((a if isinstance(a, dict) else a.__dict__).get(k)
               for k in ("email", "phone")))
    fecha = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")

    rows = [
        ("Revista", meta.get("titulo", "—")),
        ("Fuente", meta.get("url", "—")),
        ("Modo de deteccion", meta.get("modo", "—")),
        ("Fecha del analisis", fecha),
        ("Paginas de la revista", meta.get("paginas", "—")),
        ("", ""),
        ("Anunciantes detectados", n),
        ("Con sitio web", con_web),
        ("Con datos de contacto", con_contacto),
    ]
    fails = meta.get("pliegos_fallidos")
    if fails:
        paginas = ", ".join("+".join(str(x) for x in fs) for fs in fails)
        rows.append(("Paginas NO analizadas (revisar)", paginas))

    r = 5
    for label, value in rows:
        ws0.cell(row=r, column=2, value=label).font = Font(
            name="Arial", size=10, bold=True, color=XL_INK)
        ws0.cell(row=r, column=3, value=value).font = cell_font
        r += 1
    ws0.column_dimensions["B"].width = 30
    ws0.column_dimensions["C"].width = 60

    # ---------------- Hoja Anunciantes ----------------
    ws = wb.create_sheet("Anunciantes")
    ws.sheet_view.showGridLines = False
    headers = ["Marca", "Sector", "Paginas", "Web", "Email", "Telefono",
               "Tamano", "Confianza", "Metodo", "Notas"]
    widths = [26, 22, 12, 30, 30, 16, 18, 12, 12, 40]
    for c, (htext, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=htext)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 22

    for i, a in enumerate(advertisers, start=2):
        d = a if isinstance(a, dict) else a.__dict__
        web = (d.get("website") or "").strip()
        email = (d.get("email") or "").strip()
        conf = float(d.get("confidence") or 0)
        values = [
            d.get("brand", ""),
            d.get("sector", ""),
            ", ".join(str(p) for p in d.get("pages", [])),
            web, email, d.get("phone", ""),
            d.get("ad_size", ""), conf,
            d.get("method", ""), d.get("notes", ""),
        ]
        for c, val in enumerate(values, 1):
            cell = ws.cell(row=i, column=c, value=val)
            cell.font = cell_font
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 10))

        # Hipervinculos
        web_cell = ws.cell(row=i, column=4)
        if web:
            url = web if web.startswith("http") else f"https://{web}"
            web_cell.hyperlink = url
            web_cell.font = link_font
        email_cell = ws.cell(row=i, column=5)
        if email:
            email_cell.hyperlink = f"mailto:{email}"
            email_cell.font = link_font

        # Confianza como porcentaje y color segun nivel
        conf_cell = ws.cell(row=i, column=8)
        conf_cell.number_format = "0%"
        if conf >= 0.75:
            conf_cell.font = Font(name="Arial", size=10, bold=True,
                                  color=XL_GREEN)
        elif conf >= 0.5:
            conf_cell.font = Font(name="Arial", size=10, color=XL_AMBER)
        else:
            conf_cell.font = Font(name="Arial", size=10, bold=True,
                                  color=XL_RED)

        # Zebra striping
        if i % 2 == 0:
            for c in range(1, 11):
                ws.cell(row=i, column=c).fill = PatternFill(
                    "solid", fgColor=XL_GRAY_50)

    # Filtro + fila de cabecera fija
    last_row = max(1, len(advertisers) + 1)
    ws.auto_filter.ref = f"A1:J{last_row}"
    ws.freeze_panes = "A2"

    wb.save(path)


def export_json(advertisers, path: str, meta: dict):
    data = {
        "meta": meta,
        "anunciantes": [a if isinstance(a, dict) else a.__dict__
                        for a in advertisers],
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def export_html(advertisers, path: str, meta: dict):
    rows = []
    for i, a in enumerate(advertisers):
        d = a if isinstance(a, dict) else a.__dict__
        conf = d["confidence"]
        conf_class = ("alta" if conf >= 0.75 else
                      "media" if conf >= 0.5 else "baja")
        web = d["website"]
        web_link = (f'<a href="https://{web}" target="_blank">{html.escape(web)}</a>'
                    if web else '<span class="vacio">—</span>')
        email = (f'<a href="mailto:{d["email"]}">{html.escape(d["email"])}</a>'
                 if d["email"] else '<span class="vacio">—</span>')
        phone = html.escape(d["phone"]) if d["phone"] else '<span class="vacio">—</span>'
        sector = html.escape(d["sector"]) if d["sector"] else '<span class="vacio">—</span>'
        size = html.escape(d["ad_size"]) if d["ad_size"] else '<span class="vacio">—</span>'
        pages = ", ".join(str(p) for p in d["pages"]) or "—"
        rows.append(f"""
        <tr style="animation-delay:{i*0.04:.2f}s">
          <td class="marca"><span class="inicial">{html.escape(d['brand'][:1].upper())}</span>
              <span>{html.escape(d['brand'])}</span></td>
          <td>{sector}</td>
          <td class="paginas">{pages}</td>
          <td>{web_link}</td>
          <td>{email}</td>
          <td>{phone}</td>
          <td>{size}</td>
          <td><span class="conf {conf_class}">{conf:.0%}</span></td>
        </tr>""")

    n = len(advertisers)
    con_web = sum(1 for a in advertisers
                  if (a if isinstance(a, dict) else a.__dict__)["website"])
    con_contacto = sum(
        1 for a in advertisers
        if any((a if isinstance(a, dict) else a.__dict__)[k]
               for k in ("email", "phone")))
    modo = meta.get("modo", "—")
    fecha = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")

    # Reparto por sector (categorias amplias normalizadas).
    counts = Counter(
        normalize_sector((a if isinstance(a, dict) else a.__dict__)["sector"])
        for a in advertisers)
    sov_items = counts.most_common()
    sov_max = sov_items[0][1] if sov_items else 1
    sov_rows = "".join(
        f'<div class="sov-row"><div class="lab">{html.escape(sec)}</div>'
        f'<div class="track"><div class="fill" '
        f'style="width:{(c / sov_max) * 100:.1f}%"></div></div>'
        f'<div class="val">{c}</div></div>'
        for sec, c in sov_items)
    sov_block = (
        f'<div class="sov"><h2>Reparto por sector</h2>{sov_rows}</div>'
        if sov_items else "")

    doc = f"""<!doctype html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>NEVO · Anunciantes — {html.escape(meta.get('titulo','Revista'))}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Big+Shoulders:opsz,wght@10..72,700;10..72,900&family=Instrument+Sans:wght@400;500&family=Geist+Mono:wght@400;500&display=swap" rel="stylesheet">
<style>
  :root {{
    --bg:#ffffff; --ink:#0a0a0a; --orange:#ff4b00;
    --gray-50:#f4f4f2; --gray-100:#ececea; --gray-400:#9a9a97;
    --gray-800:#2a2a28;
  }}
  *{{box-sizing:border-box;margin:0;padding:0}}
  ::selection{{background:var(--orange);color:#fff}}
  body{{
    background:var(--bg); color:var(--ink);
    font-family:'Instrument Sans',ui-sans-serif,system-ui,sans-serif;
    -webkit-font-smoothing:antialiased;
    padding:56px 28px;
  }}
  .lienzo{{max-width:1240px;margin:0 auto}}

  /* Cabecera Nevo */
  .nevo-bar{{
    display:flex;align-items:center;gap:12px;margin-bottom:32px;
  }}
  .nevo-mono{{
    width:36px;height:36px;background:var(--ink);border-radius:8px;
    display:grid;place-items:center;color:#fff;
    font-family:'Big Shoulders',sans-serif;font-weight:900;font-size:24px;
    font-stretch:75%;letter-spacing:-.04em;line-height:1;
  }}
  .nevo-word{{
    font-family:'Big Shoulders',sans-serif;font-weight:900;font-size:28px;
    font-stretch:75%;letter-spacing:-.02em;line-height:1;
    text-transform:uppercase;position:relative;display:inline-block;
  }}
  .nevo-word .e{{position:relative}}
  .nevo-word .e::after{{
    content:"";position:absolute;left:0;right:.08em;bottom:-.08em;
    height:.12em;background:var(--orange);
  }}
  .nevo-tag{{
    margin-left:auto;font-family:'Geist Mono',ui-monospace,monospace;
    font-size:11px;letter-spacing:.1em;text-transform:uppercase;
    color:var(--gray-400);
  }}

  header.hero{{
    border-top:1px solid var(--gray-100);
    border-bottom:1px solid var(--gray-100);
    padding:36px 0;margin-bottom:8px;
    display:flex;justify-content:space-between;align-items:flex-end;
    flex-wrap:wrap;gap:24px;
  }}
  .kicker{{
    font-family:'Geist Mono',ui-monospace,monospace;font-weight:500;
    font-size:11px;letter-spacing:.1em;text-transform:uppercase;
    color:var(--gray-400);margin-bottom:14px;
  }}
  .kicker .dot{{color:var(--orange)}}
  h1{{
    font-family:'Big Shoulders',sans-serif;font-weight:900;
    font-stretch:75%;
    font-size:clamp(2.8rem,6vw,5rem);line-height:.92;
    letter-spacing:-.02em;text-transform:uppercase;max-width:18ch;
  }}
  .fuente{{
    text-align:right;font-family:'Geist Mono',ui-monospace,monospace;
    font-size:11px;color:var(--gray-400);line-height:1.9;
    text-transform:uppercase;letter-spacing:.06em;
  }}
  .fuente .v{{color:var(--ink);text-transform:none;letter-spacing:0}}
  .fuente a{{color:var(--ink);text-decoration:none;
            border-bottom:1px solid var(--orange);word-break:break-all;
            text-transform:none;letter-spacing:0}}

  .stats{{
    display:grid;grid-template-columns:repeat(4,1fr);gap:0;
    margin:40px 0 36px;border:1px solid var(--gray-100);
  }}
  .stat{{
    padding:24px 26px;border-right:1px solid var(--gray-100);
    background:var(--bg);position:relative;
  }}
  .stat:last-child{{border-right:none}}
  .stat .n{{
    font-family:'Big Shoulders',sans-serif;font-weight:900;
    font-stretch:75%;font-size:3.2rem;line-height:1;
    letter-spacing:-.02em;
  }}
  .stat .l{{
    font-family:'Geist Mono',ui-monospace,monospace;font-weight:500;
    font-size:11px;letter-spacing:.1em;text-transform:uppercase;
    color:var(--gray-400);margin-top:10px;
  }}
  .stat:first-child .n{{color:var(--orange)}}

  .aviso{{
    background:var(--gray-50);border-left:3px solid var(--orange);
    padding:18px 22px;margin-bottom:32px;
    font-size:14px;line-height:1.65;color:var(--gray-800);
  }}
  .aviso strong{{color:var(--ink)}}

  .sov{{margin:0 0 36px}}
  .sov h2{{
    font-family:'Big Shoulders',sans-serif;font-weight:900;font-stretch:75%;
    font-size:1.5rem;text-transform:uppercase;letter-spacing:-.01em;
    margin-bottom:16px;
  }}
  .sov-row{{
    display:grid;grid-template-columns:200px 1fr 44px;align-items:center;
    gap:14px;margin-bottom:9px;
  }}
  .sov-row .lab{{
    font-size:12.5px;color:var(--gray-800);text-align:right;
    overflow:hidden;text-overflow:ellipsis;white-space:nowrap;
  }}
  .sov-row .track{{background:var(--gray-50);height:18px;border-radius:3px;
                  overflow:hidden}}
  .sov-row .fill{{background:var(--ink);height:100%;border-radius:3px}}
  .sov-row:first-child .fill{{background:var(--orange)}}
  .sov-row .val{{
    font-family:'Geist Mono',ui-monospace,monospace;font-size:12px;
    color:var(--gray-400);
  }}
  @media(max-width:640px){{
    .sov-row{{grid-template-columns:110px 1fr 36px;gap:8px}}
  }}

  table{{
    width:100%;border-collapse:collapse;background:#fff;
    border:1px solid var(--gray-100);
  }}
  thead th{{
    background:var(--ink);color:#fff;
    font-family:'Geist Mono',ui-monospace,monospace;font-weight:500;
    font-size:10.5px;letter-spacing:.1em;text-transform:uppercase;
    padding:14px 16px;text-align:left;
  }}
  tbody tr{{
    border-bottom:1px solid var(--gray-100);
    animation:surge .5s ease both;
  }}
  tbody tr:last-child{{border-bottom:none}}
  tbody tr:hover{{background:var(--gray-50)}}
  tbody td{{padding:15px 16px;font-size:13.5px;vertical-align:middle}}
  .marca{{display:flex;align-items:center;gap:12px;font-weight:500;
         color:var(--ink)}}
  .inicial{{
    width:30px;height:30px;flex-shrink:0;border-radius:6px;
    background:var(--ink);color:#fff;display:grid;place-items:center;
    font-family:'Big Shoulders',sans-serif;font-weight:900;
    font-stretch:75%;font-size:16px;letter-spacing:-.04em;
  }}
  .paginas{{font-family:'Geist Mono',ui-monospace,monospace;
           color:var(--gray-400);font-size:12.5px}}
  td a{{color:var(--ink);text-decoration:none;
        border-bottom:1px solid var(--gray-100)}}
  td a:hover{{border-bottom-color:var(--orange)}}
  .vacio{{color:var(--gray-100)}}
  .conf{{
    font-family:'Geist Mono',ui-monospace,monospace;font-weight:500;
    font-size:11px;letter-spacing:.04em;padding:4px 9px;border-radius:3px;
    display:inline-block;
  }}
  .conf.alta{{background:var(--ink);color:#fff}}
  .conf.media{{background:var(--gray-100);color:var(--gray-800)}}
  .conf.baja{{background:var(--orange);color:#fff}}

  footer{{
    margin-top:32px;font-family:'Geist Mono',ui-monospace,monospace;
    font-size:11px;color:var(--gray-400);
    border-top:1px solid var(--gray-100);padding-top:20px;
    display:flex;justify-content:space-between;flex-wrap:wrap;gap:10px;
    text-transform:uppercase;letter-spacing:.08em;
  }}
  footer a{{color:var(--ink);text-decoration:none;
           border-bottom:1px solid var(--orange);text-transform:none;
           letter-spacing:0}}

  @keyframes surge{{from{{opacity:0;transform:translateY(6px)}}
                    to{{opacity:1;transform:translateY(0)}}}}
  @media(max-width:900px){{
    .stats{{grid-template-columns:repeat(2,1fr)}}
    .stat:nth-child(2){{border-right:none}}
    .stat:nth-child(-n+2){{border-bottom:1px solid var(--gray-100)}}
  }}
  @media(max-width:640px){{
    body{{padding:28px 14px}}
    table{{display:block;overflow-x:auto;white-space:nowrap}}
    .stats{{grid-template-columns:1fr}}
    .stat{{border-right:none;border-bottom:1px solid var(--gray-100)}}
    .stat:last-child{{border-bottom:none}}
  }}
</style>
</head>
<body>
<div class="lienzo">
  <div class="nevo-bar">
    <div class="nevo-mono">N</div>
    <div class="nevo-word">N<span class="e">E</span>VO</div>
    <div class="nevo-tag">Inteligencia competitiva · {fecha}</div>
  </div>

  <header class="hero">
    <div>
      <div class="kicker"><span class="dot">●</span> &nbsp;Informe de anunciantes</div>
      <h1>{html.escape(meta.get('titulo','Revista analizada'))}</h1>
    </div>
    <div class="fuente">
      Generado · <span class="v">{fecha}</span><br>
      Modo · <span class="v">{html.escape(modo)}</span><br>
      Fuente · <a href="{html.escape(meta.get('url',''))}">{html.escape(meta.get('url','')[:48])}</a>
    </div>
  </header>

  <div class="stats">
    <div class="stat"><div class="n">{n}</div>
      <div class="l">Anunciantes</div></div>
    <div class="stat"><div class="n">{con_web}</div>
      <div class="l">Con sitio web</div></div>
    <div class="stat"><div class="n">{con_contacto}</div>
      <div class="l">Con contacto</div></div>
    <div class="stat"><div class="n">{meta.get('paginas','—')}</div>
      <div class="l">Paginas revista</div></div>
  </div>

  <div class="aviso">
    <strong>Revision recomendada.</strong> Esta lista es una deteccion
    automatica. {"En modo gratis la separacion entre anuncio y reportaje editorial es aproximada: confirma los resultados de confianza media o baja." if modo.startswith("Heur") else "El modo IA es mas preciso, pero conviene una revision final humana antes de usar los datos comercialmente."}
  </div>

  {sov_block}

  <table>
    <thead><tr>
      <th>Marca</th><th>Sector</th><th>Paginas</th><th>Web</th>
      <th>Email</th><th>Telefono</th><th>Tamano</th><th>Confianza</th>
    </tr></thead>
    <tbody>{''.join(rows) if rows else
      '<tr><td colspan="8" style="text-align:center;padding:48px;color:var(--gray-400)">No se detectaron anunciantes.</td></tr>'}</tbody>
  </table>

  <footer>
    <span>NEVO · Extractor de Anunciantes · Deteccion {html.escape(modo)}</span>
    <span><a href="https://nevo.studio">nevo.studio</a></span>
  </footer>
</div>
</body>
</html>"""
    with open(path, "w", encoding="utf-8") as f:
        f.write(doc)


def export_all(advertisers, outdir: str, meta: dict) -> dict:
    os.makedirs(outdir, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    base = f"anunciantes_{stamp}"
    paths = {
        "html": os.path.join(outdir, base + ".html"),
        "csv": os.path.join(outdir, base + ".csv"),
        "json": os.path.join(outdir, base + ".json"),
    }
    export_html(advertisers, paths["html"], meta)
    export_csv(advertisers, paths["csv"])
    export_json(advertisers, paths["json"], meta)
    # Excel solo si la libreria esta disponible (no rompe el resto si falta).
    if OPENPYXL_AVAILABLE:
        paths["xlsx"] = os.path.join(outdir, base + ".xlsx")
        try:
            export_xlsx(advertisers, paths["xlsx"], meta)
        except Exception:  # noqa: BLE001
            paths.pop("xlsx", None)
    return paths
